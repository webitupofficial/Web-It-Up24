"""Create an empty leads.xlsx tracker with the columns send_templates.py expects.

    py create_tracker.py
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import DATA_DIR

PATH = DATA_DIR / "leads.xlsx"

COLUMNS = [
    ("Name", 22, "Contact or clinic name"),
    ("Business", 26, "Clinic / business name"),
    ("Phone", 18, "With country code, e.g. 919876543210. 10-digit numbers get 91 added."),
    ("City", 14, ""),
    ("Consent source", 30, "How they agreed to WhatsApp contact (web form, call on 05-Oct, etc). Required."),
    ("Consent date", 14, ""),
    ("Lead status", 14, "New / Verified / Do not contact. Only Verified is messaged."),
    ("Draft status", 14, "Pending / Approved. Only Approved is messaged."),
    ("Var1", 18, "Template variable {{1}}, e.g. contact name"),
    ("Var2", 18, "Template variable {{2}} (leave blank if unused)"),
    ("Var3", 18, "Template variable {{3}} (leave blank if unused)"),
    ("Notes", 30, ""),
]


def main():
    if PATH.exists():
        sys.exit(f"{PATH.name} already exists; not overwriting.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c[0] for c in COLUMNS])
    for i, (name, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F50")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    # Example test row: put your own number in Phone before testing.
    ws.append(["TEST - Me", "Webitup24", "", "", "Self (owner test)", "", "Verified", "Approved", "Raj", "", "", "Replace Phone with your own number"])

    lead_dv = DataValidation(type="list", formula1='"New,Verified,Do not contact"', allow_blank=True)
    draft_dv = DataValidation(type="list", formula1='"Pending,Approved"', allow_blank=True)
    ws.add_data_validation(lead_dv)
    ws.add_data_validation(draft_dv)
    lead_dv.add("G2:G2000")
    draft_dv.add("H2:H2000")

    help_ws = wb.create_sheet("Column guide")
    help_ws.append(["Column", "What to enter"])
    help_ws["A1"].font = help_ws["B1"].font = Font(bold=True)
    for name, _, note in COLUMNS:
        help_ws.append([name, note])
    help_ws.column_dimensions["A"].width = 16
    help_ws.column_dimensions["B"].width = 90

    wb.save(PATH)
    print(f"Created {PATH}")


if __name__ == "__main__":
    main()
