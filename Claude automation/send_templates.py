"""Send approved WhatsApp templates to verified, consenting leads.

Dry run by default. Nothing is sent unless you pass --send.

    py send_templates.py --template webitup24_clinic_intro_v1 --limit 1
    py send_templates.py --template webitup24_clinic_intro_v1 --limit 1 --send
"""
import argparse
import os
import sys
import time
from datetime import datetime

import requests
from openpyxl import load_workbook

from common import (DATA_DIR, GRAPH_API_VERSION, SEND_LOG_CSV, SEND_LOG_FIELDS, append_row,
                    load_optouts, normalize_phone, now_iso, read_rows)

REQUIRED_COLUMNS = ["Name", "Phone", "Consent source", "Lead status", "Draft status"]
VAR_COLUMNS = ["Var1", "Var2", "Var3"]


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--template", required=True, help="Approved template name in WhatsApp Manager")
    p.add_argument("--lang", default="en", help="Template language code exactly as approved (default: en)")
    p.add_argument("--tracker", default=str(DATA_DIR / "leads.xlsx"), help="Leads workbook (default: leads.xlsx)")
    p.add_argument("--sheet", default="Leads", help="Worksheet name (default: Leads)")
    p.add_argument("--limit", type=int, default=10, help="Max messages this run (default: 10)")
    p.add_argument("--daily-cap", type=int, default=100, help="Max successful sends per day (default: 100)")
    p.add_argument("--country-code", default="91",
                   help="Prefix added to 10-digit numbers without a country code (default: 91). Use '' to disable.")
    p.add_argument("--delay", type=float, default=2.0, help="Seconds between sends (default: 2)")
    p.add_argument("--retry-failed", action="store_true", help="Retry numbers whose earlier send failed")
    p.add_argument("--send", action="store_true", help="Actually send. Without this flag it is a dry run.")
    return p.parse_args()


def load_leads(path, sheet):
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"Sheet '{sheet}' not found in {path}. Sheets: {', '.join(wb.sheetnames)}")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    index = {h.lower(): i for i, h in enumerate(header) if h}
    missing = [c for c in REQUIRED_COLUMNS if c.lower() not in index]
    if missing:
        sys.exit(f"Tracker is missing columns: {', '.join(missing)}")

    leads = []
    for row_no, row in enumerate(rows, start=2):
        def cell(name):
            i = index.get(name.lower())
            v = row[i] if i is not None and i < len(row) else None
            return v
        if all(v is None for v in row):
            continue
        leads.append({
            "row": row_no,
            "name": str(cell("Name") or "").strip(),
            "phone_raw": cell("Phone"),
            "consent": str(cell("Consent source") or "").strip(),
            "lead_status": str(cell("Lead status") or "").strip().lower(),
            "draft_status": str(cell("Draft status") or "").strip().lower(),
            "vars": [str(cell(c)).strip() for c in VAR_COLUMNS if cell(c) not in (None, "")],
        })
    wb.close()
    return leads


def sent_history(template):
    """Numbers already sent / failed for this template, and successful sends today (local time)."""
    already, failed, today_count = set(), set(), 0
    today = datetime.now().date()
    for r in read_rows(SEND_LOG_CSV):
        if r.get("result") == "failed" and r.get("template") == template:
            failed.add(r.get("phone"))
        if r.get("result") != "sent":
            continue
        if r.get("template") == template:
            already.add(r.get("phone"))
        try:
            if datetime.fromisoformat(r["timestamp"]).astimezone().date() == today:
                today_count += 1
        except (KeyError, ValueError):
            pass
    return already, failed - already, today_count


def build_payload(phone, template, lang, variables):
    payload = {
        "messaging_product": "whatsapp",
        "to": phone,
        "type": "template",
        "template": {"name": template, "language": {"code": lang}},
    }
    if variables:
        payload["template"]["components"] = [{
            "type": "body",
            "parameters": [{"type": "text", "text": v} for v in variables],
        }]
    return payload


def send_one(session, url, payload):
    """Return (message_id, error)."""
    try:
        resp = session.post(url, json=payload, timeout=30)
    except requests.RequestException as e:
        return "", f"network error: {e}"
    try:
        data = resp.json()
    except ValueError:
        data = {}
    if resp.ok and data.get("messages"):
        return data["messages"][0].get("id", ""), ""
    err = data.get("error", {})
    return "", f"HTTP {resp.status_code} code={err.get('code')} {err.get('message') or resp.text[:200]}"


def main():
    args = parse_args()
    if not os.path.exists(args.tracker):
        sys.exit(f"Tracker not found: {args.tracker}. Run: py create_tracker.py")

    session = url = None
    if args.send:
        phone_id = os.environ.get("WA_PHONE_NUMBER_ID")
        token = os.environ.get("WA_ACCESS_TOKEN")
        if not phone_id or not token:
            sys.exit("Set WA_PHONE_NUMBER_ID and WA_ACCESS_TOKEN before using --send.")
        url = f"https://graph.facebook.com/{GRAPH_API_VERSION}/{phone_id}/messages"
        session = requests.Session()
        session.headers["Authorization"] = f"Bearer {token}"

    optouts = load_optouts()
    already_sent, failed_before, sent_today = sent_history(args.template)
    remaining_today = max(0, args.daily_cap - sent_today)
    budget = min(args.limit, remaining_today)

    mode = "SEND" if args.send else "DRY RUN"
    print(f"[{mode}] template={args.template} lang={args.lang} limit={args.limit} "
          f"sent_today={sent_today}/{args.daily_cap}")
    if budget == 0:
        print("Daily cap reached or limit is 0. Nothing to do.")
        return

    seen_this_run = set()
    counts = {"queued": 0, "sent": 0, "failed": 0, "skipped": 0}

    for lead in load_leads(args.tracker, args.sheet):
        if counts["queued"] >= budget:
            break
        phone = normalize_phone(lead["phone_raw"], args.country_code)
        reason = None
        if lead["lead_status"] != "verified":
            reason = "lead status not Verified"
        elif lead["draft_status"] != "approved":
            reason = "draft status not Approved"
        elif not lead["consent"]:
            reason = "no consent source recorded"
        elif not phone:
            reason = f"invalid phone '{lead['phone_raw']}'"
        elif phone in optouts:
            reason = "opted out"
        elif phone in seen_this_run:
            reason = "duplicate number in tracker"
        elif phone in already_sent:
            reason = "already sent this template"
        elif phone in failed_before and not args.retry_failed:
            reason = "failed before (see send_log.csv; use --retry-failed)"
        if reason:
            counts["skipped"] += 1
            print(f"  skip row {lead['row']:>4} {lead['name'][:25]:<25} {reason}")
            continue

        seen_this_run.add(phone)
        counts["queued"] += 1
        payload = build_payload(phone, args.template, args.lang, lead["vars"])
        label = f"row {lead['row']:>4} {lead['name'][:25]:<25} +{phone} vars={lead['vars']}"

        if not args.send:
            print(f"  would send  {label}")
            append_row(SEND_LOG_CSV, SEND_LOG_FIELDS, {
                "timestamp": now_iso(), "phone": phone, "name": lead["name"],
                "template": args.template, "result": "dry_run"})
            continue

        message_id, error = send_one(session, url, payload)
        result = "sent" if message_id else "failed"
        counts[result] += 1
        print(f"  {result:<10} {label} {message_id or error}")
        append_row(SEND_LOG_CSV, SEND_LOG_FIELDS, {
            "timestamp": now_iso(), "phone": phone, "name": lead["name"],
            "template": args.template, "result": result, "message_id": message_id, "error": error})
        if "code=190" in error:
            sys.exit("Access token rejected (code 190). Fix WA_ACCESS_TOKEN and rerun.")
        time.sleep(args.delay)

    print(f"Done. queued={counts['queued']} sent={counts['sent']} "
          f"failed={counts['failed']} skipped={counts['skipped']}")
    if not args.send and counts["queued"]:
        print("This was a dry run. Review the list above, then rerun with --send.")


if __name__ == "__main__":
    main()
